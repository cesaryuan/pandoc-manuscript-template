#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
# ]
# ///
"""
Auto-fit regular tables to window width.
This script adjusts regular tables to fit window width automatically using XML manipulation,
while skipping one-row equation layout tables generated for centered equations and labels.

Usage:
    uv run autofit_tables.py path/to/file.docx
    or as a module: autofit_tables(doc)
"""

import argparse
import re
import sys
from pathlib import Path
from typing import Optional

try:
    from docx import Document
    from docx.document import Document as DocumentObject
    from docx.table import Table
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    print("Error: python-docx is not installed. Install it with: pip install python-docx")
    sys.exit(1)


EQUATION_LAYOUT_TEXT_PATTERN = re.compile(r"^[\s\t\r\n()（）\[\]【】0-9ivxlcdmIVXLCDM.\-–—]*$")


class Colors:
    """ANSI color codes for terminal output"""
    GREEN = '\033[92m'
    CYAN = '\033[96m'
    YELLOW = '\033[93m'
    RED = '\033[91m'
    RESET = '\033[0m'


def print_success(message: str):
    """Print success message in green"""
    print(f"{Colors.GREEN}{message}{Colors.RESET}")


def print_info(message: str):
    """Print info message in cyan"""
    print(f"{Colors.CYAN}{message}{Colors.RESET}")


def print_warning(message: str):
    """Print warning message in yellow"""
    print(f"{Colors.YELLOW}{message}{Colors.RESET}")


def print_error(message: str):
    """Print error message in red"""
    print(f"{Colors.RED}{message}{Colors.RESET}")


def set_table_autofit_window(table: Table):
    """
    Set table to auto-fit to window width using XML manipulation.

    This is equivalent to Word's wdAutoFitWindow (value 2).

    Args:
        table: The table to modify
    """
    table.autofit = True
    tbl = table._element
    tblPr = tbl.tblPr

    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Set table width to 100% (percentage-based)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.insert(0, tblW)

    tblW.set(qn('w:type'), 'pct')
    tblW.set(qn('w:w'), '5000')  # 100% width (50 * 100)


def set_table_center_alignment(table: Table):
    """
    Center align the table on the page using XML manipulation.

    This is equivalent to Word's wdAlignRowCenter (value 1).

    Args:
        table: The table to modify
    """
    tbl = table._element
    tblPr = tbl.tblPr

    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Set table justification (alignment)
    jc = tblPr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        tblPr.append(jc)

    jc.set(qn('w:val'), 'center')


def table_contains_math(table: Table) -> bool:
    """Return whether a table contains Word math elements."""
    element = table._element
    return bool(element.findall(f".//{qn('m:oMath')}") or element.findall(f".//{qn('m:oMathPara')}"))


def table_non_math_text(table: Table) -> str:
    """Return visible table text outside Word math elements."""
    element = table._element
    math_elements = set(element.findall(f".//{qn('m:oMath')}") + element.findall(f".//{qn('m:oMathPara')}"))
    text_parts = []

    for text_element in element.findall(f".//{qn('w:t')}"):
        if any(parent in math_elements for parent in text_element.iterancestors()):
            continue
        if text_element.text:
            text_parts.append(text_element.text)

    return "".join(text_parts)


def is_equation_layout_table(table: Table) -> bool:
    """
    Return whether a table is used only to lay out a centered equation and label.

    Pandoc can render eqnBlockTemplate as a one-row, three-column table, such as
    an empty left cell, a centered formula cell, and a right label cell. Auto-fit
    widens that layout table and breaks the equation placement, so these tables
    are intentionally skipped.
    """
    if not table_contains_math(table):
        return False
    if len(table.rows) != 1 or len(table.columns) < 3:
        return False

    return bool(EQUATION_LAYOUT_TEXT_PATTERN.match(table_non_math_text(table)))


def autofit_tables(doc: DocumentObject, center_align: bool = True) -> int:
    """
    Auto-fit regular tables to window width.

    Args:
        doc: python-docx Document object
        center_align: Whether to center align tables (default: True)

    Returns:
        Number of tables processed
    """
    table_count = len(doc.tables)
    success_count = 0
    skipped_count = 0

    if table_count == 0:
        print_info("No tables found in document")
        return success_count

    print_info(f"Auto-fitting {table_count} table(s) to window...")

    for i, table in enumerate(doc.tables, start=1):
        try:
            if is_equation_layout_table(table):
                print_info(f"Skipping equation layout table {i}")
                skipped_count += 1
                continue

            # AutoFit to window
            set_table_autofit_window(table)

            # Optionally center align the table
            if center_align:
                set_table_center_alignment(table)

            success_count += 1

        except Exception as e:
            print_warning(f"Failed to auto-fit table {i}: {e}")

    if skipped_count:
        print_info(f"Skipped {skipped_count} equation layout table(s)")

    print_success(f"Auto-fitted {success_count} of {table_count} table(s) to window")

    return success_count


def process_file(docx_path: str, save: bool = True, center_align: bool = True) -> Optional[DocumentObject]:
    """
    Process a DOCX file to auto-fit tables.

    Args:
        docx_path: Path to the DOCX file
        save: Whether to save the document (default: True)
        center_align: Whether to center align tables (default: True)

    Returns:
        Document object if successful, None otherwise
    """
    # Validate inputs
    print_info("Validating inputs...")
    docx_file = Path(docx_path)

    if not docx_file.exists():
        print_error(f"DOCX file not found: {docx_path}")
        return None

    docx_path_abs = docx_file.resolve()
    print_info(f"Processing: {docx_path_abs}")

    try:
        # Open document
        print_info("Opening document...")
        doc = Document(str(docx_path_abs))

        # Auto-fit tables
        autofit_tables(doc, center_align=center_align)

        # Save the document
        if save:
            print_info("Saving document...")
            doc.save(str(docx_path_abs))
            print_success("Document saved")

        print_success("\nTable auto-fit processing completed successfully!")
        return doc

    except Exception as e:
        print_error(f"\nTable auto-fit processing failed: {e}")
        import traceback
        traceback.print_exc()
        return None


def main():
    """Main entry point for command-line usage"""
    parser = argparse.ArgumentParser(
        description="Auto-fit regular tables to window width in DOCX files",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  uv run autofit_tables.py manuscript.docx
  uv run autofit_tables.py output/docx/manuscript.docx
  uv run autofit_tables.py manuscript.docx --no-center

This script sets regular tables to auto-fit to window width (100%),
optionally center-aligns them on the page, and skips equation layout tables.
        """
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("--no-save", action="store_true",
                       help="Don't save the document (for testing)")
    parser.add_argument("--no-center", action="store_true",
                       help="Don't center align tables")

    args = parser.parse_args()

    result = process_file(
        args.docx_path,
        save=not args.no_save,
        center_align=not args.no_center
    )
    sys.exit(0 if result else 1)


if __name__ == "__main__":
    main()
