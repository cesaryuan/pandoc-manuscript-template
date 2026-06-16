#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
# ]
# ///
"""
Format equation layout tables in DOCX files.

Pandoc can use a one-row table to center display equations and place the
equation number in the rightmost cell. This script hides the table borders and
sets the equation-number cell's right margin to zero. It also sets the first
and last columns according to the equation number.

Usage:
    uv run format_equation_layout_tables.py path/to/file.docx
    or as a module: format_equation_layout_tables(doc)
"""

import argparse
import re
import sys
from pathlib import Path
from typing import Optional

if __package__ in (None, ""):
    # Allow direct execution via `uv run scripts/postprocess/<script>.py`.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from postprocess.autofit_tables import is_equation_layout_table, is_mathtype_marker_text

try:
    from docx.document import Document as DocumentObject
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.table import Table, _Cell
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError:
    print("Error: python-docx is not installed. Install it with: pip install python-docx")
    sys.exit(1)

from postprocess.common import (
    get_or_add_child,
    get_or_add_tbl_pr,
    open_docx,
    print_error,
    print_debug,
    print_debug_success,
    print_warning,
    save_docx,
)


BORDER_SIDES = ("top", "left", "bottom", "right", "insideH", "insideV")
CELL_BORDER_SIDES = ("top", "left", "bottom", "right")
DEFAULT_EQUATION_TABLE_WIDTH_TWIPS = 9360
SMALL_EQUATION_NUMBER_COLUMN_WIDTH_CM = 0.6
DEFAULT_EQUATION_NUMBER_COLUMN_WIDTH_CM = 0.8
TWIPS_PER_CM = 1440 / 2.54
SMALL_EQUATION_NUMBER_COLUMN_WIDTH_TWIPS = round(SMALL_EQUATION_NUMBER_COLUMN_WIDTH_CM * TWIPS_PER_CM)
DEFAULT_EQUATION_NUMBER_COLUMN_WIDTH_TWIPS = round(DEFAULT_EQUATION_NUMBER_COLUMN_WIDTH_CM * TWIPS_PER_CM)
MIN_FORMULA_COLUMN_WIDTH_TWIPS = 1440
EQUATION_NUMBER_PATTERN = re.compile(r"\d+")


def set_border_hidden(border_element) -> None:
    """Hide one Word table border element."""
    border_element.set(qn('w:val'), 'nil')
    border_element.set(qn('w:sz'), '0')
    border_element.set(qn('w:space'), '0')
    border_element.set(qn('w:color'), 'auto')


def hide_table_borders(table: Table) -> None:
    """Hide all outer and inner borders for an equation layout table."""
    tbl_pr = get_or_add_tbl_pr(table)
    tbl_borders = get_or_add_child(tbl_pr, 'w:tblBorders')
    for side in BORDER_SIDES:
        set_border_hidden(get_or_add_child(tbl_borders, f'w:{side}'))


def clear_table_style(table: Table) -> None:
    """Remove the table style before applying equation-layout formatting.

    Equation layout tables are structural helpers, so inherited table styles can
    reintroduce borders, padding, or width rules that fight the direct XML
    formatting below.
    """
    tbl_pr = get_or_add_tbl_pr(table)

    tbl_style = tbl_pr.find(qn('w:tblStyle'))
    if tbl_style is not None:
        tbl_pr.remove(tbl_style)

    tbl_look = tbl_pr.find(qn('w:tblLook'))
    if tbl_look is not None:
        tbl_pr.remove(tbl_look)


def set_table_zero_cell_margins(table: Table) -> None:
    """Set the table's default cell margins to zero on all sides."""
    tbl_pr = get_or_add_tbl_pr(table)
    tbl_cell_mar = get_or_add_child(tbl_pr, 'w:tblCellMar')

    for side in ('top', 'left', 'bottom', 'right'):
        margin = tbl_cell_mar.find(qn(f'w:{side}'))
        if margin is not None:
            tbl_cell_mar.remove(margin)

        margin = OxmlElement(f'w:{side}')
        margin.set(qn('w:w'), '0')
        margin.set(qn('w:type'), 'dxa')
        tbl_cell_mar.append(margin)


def hide_cell_borders(cell: _Cell) -> None:
    """Hide direct cell borders that can otherwise override table borders."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = get_or_add_child(tc_pr, 'w:tcBorders')

    for side in CELL_BORDER_SIDES:
        set_border_hidden(get_or_add_child(tc_borders, f'w:{side}'))


def hide_all_cell_borders(table: Table) -> None:
    """Hide borders on every cell in the equation layout table."""
    for row in table.rows:
        for cell in row.cells:
            hide_cell_borders(cell)


def set_cell_right_margin(cell: _Cell, twips: int = 0) -> None:
    """Set the right cell margin in twips; equation numbers need a zero right margin."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = get_or_add_child(tc_pr, 'w:tcMar')
    right_margin = get_or_add_child(tc_mar, 'w:right')
    right_margin.set(qn('w:w'), str(twips))
    right_margin.set(qn('w:type'), 'dxa')


def set_cell_vertical_alignment_center(cell: _Cell) -> None:
    """Center the cell content vertically inside the equation-number column."""
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = get_or_add_child(tc_pr, 'w:vAlign')
    v_align.set(qn('w:val'), 'center')


def align_equation_number_cell(cell: _Cell) -> None:
    """Make the number cell vertically centered and horizontally right-aligned."""
    set_cell_vertical_alignment_center(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def cell_visible_text(cell: _Cell) -> str:
    """Return visible text from a cell so the equation-number rule can inspect it."""
    text_parts = []
    for tag in ('w:t', 'm:t'):
        for text_element in cell._tc.findall(f".//{qn(tag)}"):
            if text_element.text:
                if tag == 'w:t' and is_mathtype_marker_text(text_element.text):
                    continue
                text_parts.append(text_element.text)
    return "".join(text_parts).strip()


def get_equation_number_value(cell: _Cell) -> int | None:
    """Return the first integer found in the equation-number cell, if present."""
    match = EQUATION_NUMBER_PATTERN.search(cell_visible_text(cell))
    if match is None:
        return None
    return int(match.group(0))


def get_equation_number_column_width(cell: _Cell) -> int:
    """Use 0.6 cm for equation numbers below 10; otherwise use 0.8 cm."""
    number = get_equation_number_value(cell)
    if number is not None and number < 10:
        return SMALL_EQUATION_NUMBER_COLUMN_WIDTH_TWIPS
    return DEFAULT_EQUATION_NUMBER_COLUMN_WIDTH_TWIPS


def set_table_fixed_layout(table: Table) -> None:
    """Use fixed table layout so Word respects the explicit equation column widths."""
    tbl_pr = get_or_add_tbl_pr(table)
    tbl_layout = get_or_add_child(tbl_pr, 'w:tblLayout')
    tbl_layout.set(qn('w:type'), 'fixed')


def get_grid_widths(table: Table) -> list[int]:
    """Return current table-grid column widths, falling back to cell widths when needed."""
    tbl_grid = table._element.find(qn('w:tblGrid'))
    widths = []

    if tbl_grid is not None:
        for grid_col in tbl_grid.findall(qn('w:gridCol')):
            width = grid_col.get(qn('w:w'))
            if width and width.isdigit():
                widths.append(int(width))

    if len(widths) == len(table.columns):
        return widths

    widths = []
    for cell in table.rows[0].cells:
        tc_w = cell._tc.tcPr.find(qn('w:tcW')) if cell._tc.tcPr is not None else None
        width = tc_w.get(qn('w:w')) if tc_w is not None else None
        widths.append(int(width) if width and width.isdigit() else 0)

    return widths


def get_or_add_table_grid(table: Table):
    """Return the table grid, creating it after table properties when absent."""
    tbl = table._element
    tbl_grid = tbl.find(qn('w:tblGrid'))
    if tbl_grid is not None:
        return tbl_grid

    tbl_grid = OxmlElement('w:tblGrid')
    insert_index = 1 if tbl.tblPr is not None else 0
    tbl.insert(insert_index, tbl_grid)
    return tbl_grid


def set_grid_widths(table: Table, widths: list[int]) -> None:
    """Write explicit table-grid column widths in twips."""
    tbl_grid = get_or_add_table_grid(table)
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)

    for width in widths:
        grid_col = OxmlElement('w:gridCol')
        grid_col.set(qn('w:w'), str(width))
        tbl_grid.append(grid_col)


def set_cell_width(cell: _Cell, width: int) -> None:
    """Set a cell width in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = get_or_add_child(tc_pr, 'w:tcW')
    tc_w.set(qn('w:w'), str(width))
    tc_w.set(qn('w:type'), 'dxa')


def set_paragraph_special_indent_none(paragraph) -> None:
    """Remove all paragraph-level first-line/hanging indent attributes.

    This mirrors Word's "Special indent: None". It must be paired with a
    paragraph style that does not itself reintroduce first-line indentation.
    """
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn('w:ind'))
    if ind is None:
        return

    for attr_name in ('w:firstLine', 'w:firstLineChars', 'w:hanging', 'w:hangingChars'):
        attr = qn(attr_name)
        if attr in ind.attrib:
            del ind.attrib[attr]

    # Drop the whole `<w:ind>` element when it only existed to carry the
    # special-indent attributes we just cleared.
    if not ind.attrib:
        p_pr.remove(ind)


def clear_paragraph_style(paragraph) -> None:
    """Remove a paragraph's explicit style so it stops inheriting Compact/a0.

    For equation layout tables we only need direct alignment/spacing, not the
    body-text style chain. Clearing `w:pStyle` is more reliable here than
    asking python-docx to switch styles by name.
    """
    p_pr = paragraph._p.get_or_add_pPr()
    p_style = p_pr.find(qn('w:pStyle'))
    if p_style is not None:
        p_pr.remove(p_style)


def normalize_equation_layout_paragraphs(table: Table) -> None:
    """Force equation/number paragraphs to use compact display-equation spacing.

    Pandoc's layout table is structural only, so the paragraphs inside it should
    not inherit body-text indentation or extra spacing that would shift the
    centered formula or the right-aligned equation number.
    """
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                clear_paragraph_style(paragraph)
                paragraph_format = paragraph.paragraph_format
                paragraph_format.space_before = Pt(2) # 非 0 防止 Word 自带公式显示不完整
                paragraph_format.space_after = Pt(2)
                paragraph_format.line_spacing = 1.0
                # After detaching from Compact/a0, remove any remaining direct
                # special-indent override so Word shows "Special: None".
                set_paragraph_special_indent_none(paragraph)


def distribute_middle_widths(current_widths: list[int], remaining_width: int) -> list[int]:
    """Distribute remaining width across formula columns while preserving proportions."""
    middle_count = max(len(current_widths) - 2, 1)
    minimum_remaining = middle_count * MIN_FORMULA_COLUMN_WIDTH_TWIPS
    remaining_width = max(remaining_width, minimum_remaining)
    middle_widths = current_widths[1:-1]
    middle_total = sum(width for width in middle_widths if width > 0)

    if middle_total <= 0:
        base_width = remaining_width // middle_count
        widths = [base_width] * middle_count
        widths[-1] += remaining_width - sum(widths)
        return widths

    distributed = [
        max(MIN_FORMULA_COLUMN_WIDTH_TWIPS, round(remaining_width * width / middle_total))
        for width in middle_widths
    ]
    if sum(distributed) > remaining_width:
        base_width = remaining_width // middle_count
        widths = [base_width] * middle_count
        widths[-1] += remaining_width - sum(widths)
        return widths

    distributed[-1] += remaining_width - sum(distributed)
    return distributed


def set_equation_layout_column_widths(table: Table, edge_width: int) -> None:
    """Set first and last columns to the selected equation-number width."""
    column_count = len(table.columns)
    if column_count < 3:
        return

    current_widths = get_grid_widths(table)
    if len(current_widths) != column_count or sum(current_widths) <= 0:
        current_widths = [0] * column_count

    total_width = sum(current_widths) or DEFAULT_EQUATION_TABLE_WIDTH_TWIPS
    middle_widths = distribute_middle_widths(
        current_widths,
        total_width - (edge_width * 2),
    )
    new_widths = [
        edge_width,
        *middle_widths,
        edge_width,
    ]

    set_table_fixed_layout(table)
    set_grid_widths(table, new_widths)

    for row in table.rows:
        for cell, width in zip(row.cells, new_widths):
            set_cell_width(cell, width)


def get_equation_number_cell(table: Table) -> _Cell:
    """Return the rightmost cell, which holds the equation number in the layout template."""
    return table.rows[0].cells[-1]


def format_equation_layout_table(table: Table) -> None:
    """Apply border, margin, and column-width fixes to one equation layout table."""
    number_cell = get_equation_number_cell(table)
    edge_width = get_equation_number_column_width(number_cell)

    clear_table_style(table)
    set_table_zero_cell_margins(table)
    hide_table_borders(table)
    hide_all_cell_borders(table)
    set_cell_right_margin(number_cell, 0)
    set_equation_layout_column_widths(table, edge_width)
    normalize_equation_layout_paragraphs(table)
    align_equation_number_cell(number_cell)


def format_equation_layout_tables(doc: DocumentObject) -> int:
    """
    Format all detected equation layout tables.

    Args:
        doc: python-docx Document object

    Returns:
        Number of equation layout tables formatted
    """
    table_count = len(doc.tables)
    processed_count = 0

    if table_count == 0:
        print_debug("No tables found in document")
        return processed_count

    print_debug(f"Formatting equation layout tables among {table_count} table(s)...")
    for i, table in enumerate(doc.tables, start=1):
        try:
            if not is_equation_layout_table(table):
                continue

            format_equation_layout_table(table)
            processed_count += 1
            print_debug(f"Formatted equation layout table {i}")
        except Exception as e:
            print_warning(f"Failed to format equation layout table {i}: {e}")

    print_debug_success(f"Formatted {processed_count} equation layout table(s)")
    return processed_count


def process_file(docx_path: str, save: bool = True) -> Optional[DocumentObject]:
    """
    Process a DOCX file to format equation layout tables.

    Args:
        docx_path: Path to the DOCX file
        save: Whether to save the document (default: True)

    Returns:
        Document object if successful, None otherwise
    """
    try:
        doc, docx_path_abs = open_docx(docx_path)
        if doc is None or docx_path_abs is None:
            return None

        processed_count = format_equation_layout_tables(doc)

        if save:
            if processed_count > 0:
                save_docx(doc, docx_path_abs)
            else:
                print_debug("No changes made, skipping save")

        print_debug_success("\nEquation layout table formatting completed successfully!")
        return doc

    except Exception as e:
        print_error(f"\nEquation layout table formatting failed: {e}")
        import traceback
        traceback.print_exc()
        return None


def main():
    """Main entry point for command-line usage"""
    parser = argparse.ArgumentParser(
        description="Format equation layout tables in DOCX files",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  uv run format_equation_layout_tables.py manuscript.docx
  uv run format_equation_layout_tables.py output/docx/manuscript.docx
  uv run format_equation_layout_tables.py manuscript.docx --no-save

This script hides all borders on equation layout tables, sets the right margin
of the equation-number cell to zero, and sets the first and last columns to
0.6 cm for equation numbers below 10, otherwise 0.8 cm.
        """
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("--no-save", action="store_true",
                       help="Don't save the document")

    args = parser.parse_args()

    result = process_file(
        args.docx_path,
        save=not args.no_save
    )
    sys.exit(0 if result else 1)


if __name__ == "__main__":
    main()
