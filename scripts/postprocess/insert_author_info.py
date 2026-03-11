#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
#   "pyyaml>=6.0",
#   "lxml>=4.9.0",
# ]
# ///
"""
Insert author information into DOCX file after the title.

This script reads author metadata from the manuscript YAML header and inserts
formatted author names, affiliations, and corresponding author footnote.

Usage:
    uv run insert_author_info.py manuscript.docx manuscript.md

Supported YAML structures in manuscript.md:
    authors:
      - name: Author 1
        affiliations: [a, b, c]
      - name: Author 2
        affiliations: [a, b, c]
        corresponding: true
        email: author2@university.edu
        title: Professor

    affiliations:
      a: School of XXXX Department, XXXX University, XXXX, China
      b: Key Laboratory of XXXX, XXXX University, XXXX, China
      c: Key Laboratory of XXXX, XXXX University, XXXX, China

Or inline affiliation text per author:
    authors:
      - name: Author 1
        affiliations:
          - School of XXXX Department, XXXX University, XXXX, China
          - Key Laboratory of XXXX, XXXX University, XXXX, China
      - name: Author 2
        affiliations:
          - School of XXXX Department, XXXX University, XXXX, China
          - Key Laboratory of XXXX, XXXX University, XXXX, China
        corresponding: true
        email: author2@university.edu
        title: Professor
"""

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from lxml import etree
except ImportError as e:
    print(f"Error: Missing dependency: {e}")
    print("Install with: pip install python-docx pyyaml lxml")
    sys.exit(1)

try:
    import yaml
except ImportError:
    print("Error: pyyaml is not installed. Install it with: pip install pyyaml")
    sys.exit(1)


# Word XML namespaces
WORD_NAMESPACE = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'


def parse_yaml_header(md_path: str) -> dict:
    """Parse YAML front matter from a markdown file."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    match = re.match(r'^---\s*\n(.*?)\n---', content, re.DOTALL)
    if not match:
        raise ValueError("No YAML front matter found in markdown file")

    return yaml.safe_load(match.group(1))


def normalize_author_metadata(metadata: dict) -> tuple[list[dict], dict[str, str]]:
    """Normalize author and affiliation metadata to keyed affiliation form."""
    authors = metadata.get('authors', metadata.get('author', [])) or []
    raw_affiliations = metadata.get('affiliations', metadata.get('affiliation', {})) or {}

    normalized_affiliations: dict[str, str] = {}
    affiliation_index_by_text: dict[str, str] = {}

    def index_to_alpha(index: int) -> str:
        """Convert 1-based index to alphabetical labels: 1->a, 26->z, 27->aa."""
        label = ''
        current = index
        while current > 0:
            current -= 1
            label = chr(ord('a') + (current % 26)) + label
            current //= 26
        return label

    def register_affiliation(text: str) -> str:
        cleaned_text = str(text).strip()
        if not cleaned_text:
            return ''

        existing_key = affiliation_index_by_text.get(cleaned_text)
        if existing_key:
            return existing_key

        new_key = index_to_alpha(len(normalized_affiliations) + 1)
        normalized_affiliations[new_key] = cleaned_text
        affiliation_index_by_text[cleaned_text] = new_key
        return new_key

    if isinstance(raw_affiliations, dict):
        for key, value in raw_affiliations.items():
            normalized_key = str(key)
            normalized_value = str(value).strip()
            if not normalized_value:
                continue
            normalized_affiliations[normalized_key] = normalized_value
            affiliation_index_by_text[normalized_value] = normalized_key

    normalized_authors = []
    for author in authors:
        if not isinstance(author, dict):
            continue

        normalized_author = dict(author)
        raw_author_affiliations = author.get('affiliations', author.get('affiliation', [])) or []

        if isinstance(raw_author_affiliations, str):
            raw_author_affiliations = [raw_author_affiliations]

        normalized_keys = []
        for affiliation in raw_author_affiliations:
            if affiliation is None:
                continue

            if isinstance(affiliation, str):
                cleaned_affiliation = affiliation.strip()
            else:
                cleaned_affiliation = str(affiliation).strip()

            if not cleaned_affiliation:
                continue

            if cleaned_affiliation in normalized_affiliations:
                normalized_keys.append(cleaned_affiliation)
                continue

            if cleaned_affiliation in affiliation_index_by_text:
                normalized_keys.append(affiliation_index_by_text[cleaned_affiliation])
                continue

            normalized_keys.append(register_affiliation(cleaned_affiliation))

        normalized_author['affiliations'] = normalized_keys
        normalized_authors.append(normalized_author)

    return normalized_authors, normalized_affiliations


# =============================================================================
# XML Element Helpers (for low-level OOXML operations)
# =============================================================================

def create_element(tag: str, **attribs) -> etree._Element:
    """Create an OOXML element with optional attributes."""
    element = etree.Element(qn(tag), nsmap={'w': WORD_NAMESPACE})
    for key, value in attribs.items():
        element.set(qn(f'w:{key}'), str(value))
    return element


def create_text_element(text: str) -> etree._Element:
    """Create a w:t element with space preservation."""
    t = create_element('w:t')
    t.text = text
    t.set(XML_SPACE, 'preserve')
    return t


def set_font_properties(rPr, font_name: str, font_size: int | None = None,
                        superscript: bool = False, italic: bool = False):
    """Set font properties on a run properties element."""
    # Font family
    rFonts = create_element('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)

    # Font size (in half-points)
    if font_size:
        rPr.append(create_element('w:sz', val=str(font_size)))
        rPr.append(create_element('w:szCs', val=str(font_size)))

    if superscript:
        rPr.append(create_element('w:vertAlign', val='superscript'))

    if italic:
        rPr.append(create_element('w:i'))


def add_run_with_text(paragraph_elem, text: str, font_name: str = "Times New Roman",
                      font_size: int = 24, superscript: bool = False, italic: bool = False):
    """Add a run with text to a paragraph element."""
    run = create_element('w:r')
    rPr = create_element('w:rPr')
    set_font_properties(rPr, font_name, font_size, superscript, italic)
    run.append(rPr)
    run.append(create_text_element(text))
    paragraph_elem.append(run)
    return run


def create_centered_paragraph() -> etree._Element:
    """Create a centered paragraph element."""
    p = create_element('w:p')
    pPr = create_element('w:pPr')
    pPr.append(create_element('w:jc', val='center'))
    p.append(pPr)
    return p



# =============================================================================
# Footnote Handling (requires low-level XML - python-docx doesn't support custom marks)
# =============================================================================

FOOTNOTES_XML_TEMPLATE = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
             xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
    <w:footnote w:type="separator" w:id="-1">
        <w:p><w:r><w:separator/></w:r></w:p>
    </w:footnote>
    <w:footnote w:type="continuationSeparator" w:id="0">
        <w:p><w:r><w:continuationSeparator/></w:r></w:p>
    </w:footnote>
</w:footnotes>'''


def get_or_create_footnotes_part(doc):
    """Get or create the footnotes part in the document."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI

    document_part = doc.part

    # Check if footnotes part exists
    for rel in document_part.rels.values():
        if rel.reltype == RT.FOOTNOTES:
            return rel.target_part

    # Create new footnotes part
    footnotes_part = Part(
        PackURI('/word/footnotes.xml'),
        'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',
        FOOTNOTES_XML_TEMPLATE.encode('utf-8'),
        document_part.package
    )
    document_part.relate_to(footnotes_part, RT.FOOTNOTES)
    return footnotes_part


def add_footnote_with_custom_mark(doc, run_element, footnote_text: str, footnote_id: int = 1,
                                  font_name: str = "Times New Roman", font_size: int = 20,
                                  custom_mark: str = "*"):
    """Add a Word footnote with custom mark (*) to the document."""
    footnotes_part = get_or_create_footnotes_part(doc)
    footnotes_xml = etree.fromstring(footnotes_part.blob)

    # Create footnote content
    footnote = create_element('w:footnote', id=str(footnote_id))
    p = etree.SubElement(footnote, qn('w:p'))

    # Paragraph style
    pPr = etree.SubElement(p, qn('w:pPr'))
    pStyle = etree.SubElement(pPr, qn('w:pStyle'))
    pStyle.set(qn('w:val'), 'FootnoteText')

    # Add italic text run
    add_run_with_text(p, footnote_text, font_name, font_size, italic=True)

    footnotes_xml.append(footnote)
    footnotes_part._blob = etree.tostring(footnotes_xml, xml_declaration=True, encoding='UTF-8', standalone='yes')

    # Add footnote reference in main document
    rPr = create_element('w:rPr')
    set_font_properties(rPr, font_name, superscript=True)
    run_element.insert(0, rPr)

    # Create footnoteReference with customMarkFollows
    footnote_ref = create_element('w:footnoteReference', customMarkFollows='1', id=str(footnote_id))
    run_element.append(footnote_ref)
    run_element.append(create_text_element(custom_mark))


# =============================================================================
# Author/Affiliation Paragraph Creation
# =============================================================================

def create_author_paragraph(authors: list, font_name: str = "Times New Roman", font_size: int = 24):
    """Create a paragraph with author names and superscript affiliations."""
    p = create_centered_paragraph()
    footnote_run = None

    for i, author in enumerate(authors):
        add_run_with_text(p, author['name'], font_name, font_size)

        # Add superscript affiliations
        if affiliations := author.get('affiliations', author.get('affiliation', [])):
            add_run_with_text(p, ','.join(affiliations), font_name, font_size, superscript=True)

            # Mark corresponding author for footnote
            if author.get('corresponding', False):
                footnote_run = create_element('w:r')
                rPr = create_element('w:rPr')
                set_font_properties(rPr, font_name, superscript=True)
                footnote_run.append(rPr)
                p.append(footnote_run)

        # Add comma separator (except for last author)
        if i < len(authors) - 1:
            add_run_with_text(p, ', ', font_name, font_size)

    return p, footnote_run


def create_affiliation_paragraph(key: str, text: str, font_name: str = "Times New Roman", font_size: int = 20):
    """Create a paragraph for a single affiliation."""
    p = create_centered_paragraph()
    add_run_with_text(p, key, font_name, font_size, superscript=True, italic=True)
    add_run_with_text(p, ' ' + text, font_name, font_size, italic=True)
    return p


def find_title_paragraph_index(doc) -> int:
    """Find the index of the title paragraph in the document."""
    for i, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else ''
        if 'Title' in style_name:
            return i
        if para.text.strip() and i == 0:
            return i
    return 0


def insert_author_info_to_doc(doc, md_path: str) -> tuple[int, int, bool]:
    """
    Insert author information into a Document object after the title.

    Args:
        doc: The Document object to modify
        md_path: Path to the markdown file with YAML metadata

    Returns:
        Tuple of (authors_count, affiliations_count, has_footnote)
    """
    # Parse YAML metadata
    metadata = parse_yaml_header(md_path)

    authors, affiliations = normalize_author_metadata(metadata)

    if not authors:
        return (0, 0, False)

    # Find title paragraph
    title_idx = find_title_paragraph_index(doc)
    title_para = doc.paragraphs[title_idx]

    # Find corresponding author and build footnote text
    footnote_text = ''
    for author in authors:
        if author.get('corresponding', False):
            name = author.get('name', '')
            title = author.get('title', '')
            email = author.get('email', '')
            # Get first affiliation text
            affiliation_text = ''
            if author.get('affiliations') and affiliations:
                first_affil_key = author['affiliations'][0]
                affiliation_text = affiliations.get(first_affil_key, '')

            footnote_text = f'*Correspondence to: Dr. {name}'
            if title:
                footnote_text += f' ({title})'
            if affiliation_text:
                footnote_text += f', {affiliation_text}'
            if email:
                footnote_text += f'. E-mail: {email}.'
            break

    # Create author paragraph with footnote reference
    author_para_elem, footnote_run = create_author_paragraph(authors)

    # Insert author paragraph after title
    title_para._element.addnext(author_para_elem)
    current_elem = author_para_elem

    # Create and insert affiliation paragraphs
    for key in sorted(affiliations.keys(), key=lambda value: (not str(value).isdigit(), int(value) if str(value).isdigit() else str(value))):
        affil_para_elem = create_affiliation_paragraph(key, affiliations[key])
        current_elem.addnext(affil_para_elem)
        current_elem = affil_para_elem

    # Add footnote if there's a corresponding author
    has_footnote = False
    if footnote_run is not None and footnote_text:
        # Add Word footnote with custom mark (*)
        add_footnote_with_custom_mark(doc, footnote_run, footnote_text, footnote_id=1)
        has_footnote = True

    return (len(authors), len(affiliations), has_footnote)


def insert_author_info(docx_path: str, md_path: str) -> tuple[int, int, bool]:
    """
    Insert author information into a DOCX file after the title.

    Args:
        docx_path: Path to the DOCX file
        md_path: Path to the markdown file with YAML metadata

    Returns:
        Tuple of (authors_count, affiliations_count, has_footnote)
    """
    doc = Document(docx_path)
    result = insert_author_info_to_doc(doc, md_path)

    if result[0] == 0:
        print("Warning: No authors found in YAML metadata")
        return result

    doc.save(docx_path)
    return result


def main():
    """Main entry point for command-line usage."""
    parser = argparse.ArgumentParser(
        description="Insert author information into DOCX file after the title",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  uv run insert_author_info.py manuscript.docx manuscript.md
  uv run insert_author_info.py output/manuscript.docx manuscript.md

Supported YAML structures in manuscript.md:
    authors:
      - name: Author 1
        affiliations: [a, b, c]
      - name: Author 2
        affiliations: [a, b, c]
        corresponding: true
        email: author2@university.edu
        title: Professor
    affiliations:
      a: School of XXXX Department, XXXX University, XXXX, China
      b: Key Laboratory of XXXX, XXXX University, XXXX, China
      c: Key Laboratory of XXXX, XXXX University, XXXX, China

Or inline affiliation text per author:
    authors:
      - name: Author 1
        affiliations:
          - School of XXXX Department, XXXX University, XXXX, China
          - Key Laboratory of XXXX, XXXX University, XXXX, China
      - name: Author 2
        affiliations:
          - School of XXXX Department, XXXX University, XXXX, China
          - Key Laboratory of XXXX, XXXX University, XXXX, China
        corresponding: true
        email: author2@university.edu
        title: Professor
        """
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("md_path", help="Path to the markdown file with YAML metadata")

    args = parser.parse_args()

    if not Path(args.docx_path).exists():
        print(f"Error: DOCX file not found: {args.docx_path}")
        sys.exit(1)

    if not Path(args.md_path).exists():
        print(f"Error: Markdown file not found: {args.md_path}")
        sys.exit(1)

    try:
        authors_count, affiliations_count, has_footnote = insert_author_info(
            args.docx_path, args.md_path
        )
        print(f"Successfully inserted author information:")
        print(f"  - Authors: {authors_count}")
        print(f"  - Affiliations: {affiliations_count}")
        print(f"  - Corresponding author footnote: {'Yes' if has_footnote else 'No'}")
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
