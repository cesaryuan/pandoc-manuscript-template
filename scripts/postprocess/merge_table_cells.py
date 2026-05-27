#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
# ]
# ///
"""
Merge table cells based on markers (<< or !<! for left merge, ^^ or !^! for up merge)
This script implements the same functionality as merge-table-cells.ps1

Usage:
    uv run merge_table_cells.py path/to/file.docx
    or as a module: merge_table_cells(doc)
"""

import argparse
import sys
from pathlib import Path
from typing import Optional

try:
    from docx import Document
    from docx.document import Document as DocumentObject
    from docx.table import _Cell
except ImportError:
    print("Error: python-docx is not installed. Install it with: pip install python-docx")
    sys.exit(1)


class Colors:
    """ANSI color codes for terminal output"""
    GREEN = '\033[92m'
    CYAN = '\033[96m'
    YELLOW = '\033[93m'
    RED = '\033[91m'
    RESET = '\033[0m'


# Keep legacy markers so existing manuscripts continue to build.
LEFT_MERGE_MARKERS = {"!<!", "<<"}
UP_MERGE_MARKERS = {"!^!", "^^"}


def print_success(message: str):
    """Print success message in green"""
    print(f"{Colors.GREEN}{message}{Colors.RESET}")


def print_info(message: str):
    """Print info message in cyan"""
    print(f"{Colors.CYAN}{message}{Colors.RESET}")


def print_warning(message: str):
    """Print warning message in yellow"""
    print(f"{Colors.YELLOW}{message}{Colors.RESET}")


def print_error(message: str):
    """Print error message in red"""
    print(f"{Colors.RED}{message}{Colors.RESET}")


def get_clean_cell_text(cell: _Cell) -> str:
    """
    Clean cell text by removing trailing hidden characters.
    Word cells contain trailing paragraph markers and other characters.
    """
    text = cell.text.strip()
    return text


def merge_table_cells(doc: DocumentObject) -> tuple[int, int]:
    """
    Merge table cells based on special markers.

    Markers:
        << or !<! - Merge with cell to the left
        ^^ or !^! - Merge with cell above

    Args:
        doc: python-docx Document object

    Returns:
        Tuple of (left_merge_count, up_merge_count)
    """
    left_merge_count = 0
    up_merge_count = 0

    table_count = len(doc.tables)

    if table_count == 0:
        print_info("No tables found in document")
        return left_merge_count, up_merge_count

    print_info(f"Processing {table_count} table(s)...")

    for table in doc.tables:
        # ===================================================================
        # Phase 1: Process left merges (<< / !<!)
        # Iterate from last row to first, last column to second column
        # ===================================================================

        for row_idx in range(len(table.rows) - 1, -1, -1):
            row = table.rows[row_idx]

            # Iterate from last cell to second cell (reverse order)
            for col_idx in range(len(row.cells) - 1, 0, -1):
                try:
                    cell = row.cells[col_idx]
                    cell_text = get_clean_cell_text(cell)

                    if cell_text in LEFT_MERGE_MARKERS:
                        # Clear the marker text
                        cell.text = ""

                        # Merge with left cell
                        try:
                            left_cell = row.cells[col_idx - 1]
                            merged_cell = left_cell.merge(cell)

                            # Remove empty paragraphs from merged cell
                            paragraphs_to_remove = []
                            for paragraph in merged_cell.paragraphs:
                                if not paragraph.text.strip():
                                    paragraphs_to_remove.append(paragraph)

                            for paragraph in paragraphs_to_remove:
                                p_element = paragraph._element
                                p_element.getparent().remove(p_element)

                            left_merge_count += 1
                        except Exception as e:
                            print_warning(f"Failed to merge left at row {row_idx + 1}, cell {col_idx + 1}: {e}")

                except Exception:
                    # Skip cells that cause errors (might be part of merged cells)
                    continue

        # ===================================================================
        # Phase 2: Process up merges (^^ / !^!)
        # Iterate from last row to second row
        # ===================================================================

        for row_idx in range(len(table.rows) - 1, 0, -1):
            row = table.rows[row_idx]

            # Check each cell in current row
            for col_idx in range(len(row.cells) - 1, -1, -1):
                try:
                    cell = row.cells[col_idx]
                    cell_text = get_clean_cell_text(cell)

                    if cell_text in UP_MERGE_MARKERS:
                        # Clear the marker text
                        cell.text = ""

                        # Merge with cell above
                        try:
                            above_cell = table.rows[row_idx - 1].cells[col_idx]
                            merged_cell = above_cell.merge(cell)

                            # Remove empty paragraphs from merged cell
                            paragraphs_to_remove = []
                            for paragraph in merged_cell.paragraphs:
                                if not paragraph.text.strip():
                                    paragraphs_to_remove.append(paragraph)

                            for paragraph in paragraphs_to_remove:
                                p_element = paragraph._element
                                p_element.getparent().remove(p_element)

                            up_merge_count += 1
                        except Exception as e:
                            print_warning(f"Failed to merge up at row {row_idx + 1}, cell {col_idx + 1}: {e}")

                except Exception:
                    # Skip cells that cause errors (might be part of merged cells)
                    continue

    return left_merge_count, up_merge_count


def process_file(docx_path: str, save: bool = True) -> Optional[DocumentObject]:
    """
    Process a DOCX file to merge table cells.

    Args:
        docx_path: Path to the DOCX file
        save: Whether to save the document (default: True)

    Returns:
        Document object if successful, None otherwise
    """
    # Validate inputs
    print_info("Validating inputs...")
    docx_file = Path(docx_path)

    if not docx_file.exists():
        print_error(f"DOCX file not found: {docx_path}")
        return None

    docx_path_abs = docx_file.resolve()
    print_info(f"Processing: {docx_path_abs}")

    try:
        # Open document
        print_info("Opening document...")
        doc = Document(str(docx_path_abs))

        # Process table cell merges
        print_info("Processing table cell merges...")
        left_merges, up_merges = merge_table_cells(doc)

        print_success(f"Processed {len(doc.tables)} table(s)")
        print_success(f"Left merges: {left_merges}")
        print_success(f"Up merges: {up_merges}")

        # Save the document
        if save:
            print_info("Saving document...")
            doc.save(str(docx_path_abs))
            print_success("Document saved")

        print_success("\nTable cell merge processing completed successfully!")
        return doc

    except Exception as e:
        print_error(f"\nTable cell merge processing failed: {e}")
        import traceback
        traceback.print_exc()
        return None


def main():
    """Main entry point for command-line usage"""
    parser = argparse.ArgumentParser(
        description="Merge table cells in DOCX files based on markers (<<, !<!, ^^, and !^!)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  uv run merge_table_cells.py manuscript.docx
  uv run merge_table_cells.py output/docx/manuscript.docx

Markers:
  << or !<!  - Merge with cell to the left
  ^^ or !^!  - Merge with cell above
        """
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("--no-save", action="store_true",
                       help="Don't save the document (for testing)")

    args = parser.parse_args()

    result = process_file(args.docx_path, save=not args.no_save)
    sys.exit(0 if result else 1)


if __name__ == "__main__":
    main()
