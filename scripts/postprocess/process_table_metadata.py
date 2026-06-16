#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
# ]
# ///
"""
Process table metadata from captions.
This script parses metadata in table captions (format: |key=value key2=value2|)
and applies the settings to tables, then removes the metadata from captions.

Usage:
    uv run process_table_metadata.py path/to/file.docx
    or as a module: process_table_metadata(doc)
"""

import argparse
import re
import sys
from pathlib import Path
from typing import Optional, Dict

if __package__ in (None, ""):
    # Allow direct execution via `uv run scripts/postprocess/<script>.py`.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx.document import Document as DocumentObject
from docx.table import Table
from docx.shared import Pt, Cm, Mm, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from postprocess.common import (
    get_or_add_tbl_pr,
    open_docx,
    print_error,
    print_debug,
    print_debug_success,
    print_warning,
    save_docx,
)


def parse_table_metadata(caption_text: str) -> Dict[str, str]:
    """
    Parse metadata from caption text.
    Format: |key=value key2=value2|

    Args:
        caption_text: The caption text containing metadata

    Returns:
        Dictionary of key-value pairs
    """
    metadata = {}

    # Match pattern: |key=value key2=value2|
    match = re.search(r'\|([^|]+)\|\s*$', caption_text)
    if match:
        metadata_string = match.group(1).strip()

        # Split by spaces and parse key=value pairs
        pairs = metadata_string.split()
        for pair in pairs:
            kv_match = re.match(r'^([^=]+)=(.+)$', pair)
            if kv_match:
                key = kv_match.group(1).strip()
                value = kv_match.group(2).strip()
                metadata[key] = value

    return metadata


def remove_metadata_from_caption(caption_text: str) -> str:
    """
    Remove metadata from caption text.

    Args:
        caption_text: The caption text containing metadata

    Returns:
        Clean caption text without metadata
    """
    # Remove the |...| pattern and trim whitespace
    clean_text = re.sub(r'\s*\|[^|]+\|\s*', '', caption_text)
    return clean_text.strip()


def convert_to_points(dimension: str) -> float:
    """
    Convert dimension string to points (Word uses points internally).

    Args:
        dimension: Dimension with unit (e.g., "0.10cm", "5pt", "0.5in")

    Returns:
        Value in points
    """
    # Parse dimension with unit
    match = re.match(r'^([\d.]+)(cm|mm|in|pt)$', dimension.lower())
    if match:
        value = float(match.group(1))
        unit = match.group(2)

        if unit == 'cm':
            return Cm(value).pt
        elif unit == 'mm':
            return Mm(value).pt
        elif unit == 'in':
            return Inches(value).pt
        elif unit == 'pt':
            return Pt(value).pt

    # Default: assume points if no unit specified
    return float(dimension)


def set_cell_margins(table: Table, top=None, bottom=None, left=None, right=None):
    """
    Set cell margins for a table using XML manipulation.

    Args:
        table: The table to modify
        top, bottom, left, right: Margin values in points (optional)
    """
    tblPr = get_or_add_tbl_pr(table)

    # Find or create tblCellMar element
    tblCellMar = tblPr.find(qn('w:tblCellMar'))
    if tblCellMar is None:
        tblCellMar = OxmlElement('w:tblCellMar')
        tblPr.append(tblCellMar)

    # Set margins
    margins = {
        'top': top,
        'bottom': bottom,
        'left': left,
        'right': right
    }

    for side, value in margins.items():
        if value is not None:
            # Remove existing margin element
            existing = tblCellMar.find(qn(f'w:{side}'))
            if existing is not None:
                tblCellMar.remove(existing)

            # Create new margin element
            margin_elem = OxmlElement(f'w:{side}')
            margin_elem.set(qn('w:w'), str(int(value * 20)))  # Convert to twips (1/20 of a point)
            margin_elem.set(qn('w:type'), 'dxa')
            tblCellMar.append(margin_elem)


def set_table_alignment(table: Table, alignment: str):
    """
    Set table alignment using XML manipulation.

    Args:
        table: The table to modify
        alignment: 'left', 'center', or 'right'
    """
    tblPr = get_or_add_tbl_pr(table)

    # Find or create jc (justification) element
    jc = tblPr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        tblPr.append(jc)

    # Set alignment
    alignment_map = {
        'left': 'left',
        'center': 'center',
        'right': 'right'
    }
    jc.set(qn('w:val'), alignment_map.get(alignment.lower(), 'center'))


def set_autofit_behavior(table: Table, behavior: str):
    """
    Set table autofit behavior using XML manipulation.

    Args:
        table: The table to modify
        behavior: 'fixed', 'content', or 'window'
    """
    tblPr = get_or_add_tbl_pr(table)

    # Set table width based on behavior
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.insert(0, tblW)

    if behavior.lower() == 'window':
        # Auto width (percentage-based)
        tblW.set(qn('w:type'), 'pct')
        tblW.set(qn('w:w'), '5000')  # 100% (50 * 100)
    elif behavior.lower() == 'content':
        # Auto width (fit content)
        tblW.set(qn('w:type'), 'auto')
        tblW.set(qn('w:w'), '0')
    elif behavior.lower() == 'fixed':
        # Fixed width - preserve current width
        pass


def apply_table_metadata(table: Table, metadata: Dict[str, str]) -> list[str]:
    """
    Apply metadata settings to table.

    Args:
        table: The table to modify
        metadata: Dictionary of metadata key-value pairs

    Returns:
        List of applied settings descriptions
    """
    applied_settings = []

    for key, value in metadata.items():
        try:
            key_lower = key.lower()

            if key_lower == 'cell_margin':
                # Set cell margins (all sides)
                points = convert_to_points(value)
                set_cell_margins(table, top=points, bottom=points, left=points, right=points)
                applied_settings.append(f"cell_margin={value}")

            elif key_lower == 'cell_margin_top':
                points = convert_to_points(value)
                set_cell_margins(table, top=points)
                applied_settings.append(f"cell_margin_top={value}")

            elif key_lower == 'cell_margin_bottom':
                points = convert_to_points(value)
                set_cell_margins(table, bottom=points)
                applied_settings.append(f"cell_margin_bottom={value}")

            elif key_lower == 'cell_margin_left':
                points = convert_to_points(value)
                set_cell_margins(table, left=points)
                applied_settings.append(f"cell_margin_left={value}")

            elif key_lower == 'cell_margin_right':
                points = convert_to_points(value)
                set_cell_margins(table, right=points)
                applied_settings.append(f"cell_margin_right={value}")

            elif key_lower == 'cell_spacing':
                # Note: python-docx doesn't directly support cell spacing
                # This would require advanced XML manipulation
                print_warning("cell_spacing is not fully supported in python-docx")
                applied_settings.append(f"cell_spacing={value} (limited support)")

            elif key_lower == 'row_height':
                # Set row height for all rows
                points = convert_to_points(value)
                for row in table.rows:
                    row.height = Pt(points)
                applied_settings.append(f"row_height={value}")

            elif key_lower == 'alignment':
                # Set table alignment
                set_table_alignment(table, value)
                applied_settings.append(f"alignment={value}")

            elif key_lower == 'autofit':
                # Set autofit behavior
                set_autofit_behavior(table, value)
                applied_settings.append(f"autofit={value}")

            else:
                print_warning(f"Unknown metadata key: {key}")

        except Exception as e:
            print_warning(f"Failed to apply setting {key}={value}: {e}")

    return applied_settings


def process_table_metadata(doc: DocumentObject) -> tuple[int, int]:
    """
    Process table metadata from captions.

    Args:
        doc: python-docx Document object

    Returns:
        Tuple of (processed_count, total_settings_applied)
    """
    processed_count = 0
    total_settings_applied = 0

    table_count = len(doc.tables)

    if table_count == 0:
        print_debug("No tables found in document")
        return processed_count, total_settings_applied

    print_debug(f"Processing {table_count} table(s)...")

    # Iterate through all tables
    for i, table in enumerate(doc.tables, start=1):
        # Find the table's caption (can be before or after the table)
        caption_para = None
        caption_text = ""

        # Get the previous and next elements
        prev_element = table._element.getprevious()
        next_element = table._element.getnext()

        # Function to check if an element is a caption with metadata
        def check_caption(element):
            if element is None or not element.tag.endswith('p'):
                return None, ""

            # Find the corresponding paragraph object
            for para in doc.paragraphs:
                if para._element == element:
                    style_name = para.style.name if para.style else ""
                    # Check if it's a caption style (Table Caption, Caption, 题注，etc.)
                    if style_name and ('Caption' in style_name or '题注' in style_name or 'Table' in style_name):
                        # Check if this caption contains metadata
                        if re.search(r'\|[^|]+=[^|]+\|', para.text):
                            return para, para.text
            return None, ""

        # First check previous element (Pandoc usually puts caption before table)
        caption_para, caption_text = check_caption(prev_element)

        # If not found, check next element
        if not caption_para:
            caption_para, caption_text = check_caption(next_element)

        # Process metadata if caption found
        if caption_para and re.search(r'\|[^|]+=[^|]+\|', caption_text):
            print_debug(f"Processing Table {i}...")

            # Parse metadata
            metadata = parse_table_metadata(caption_text)

            if metadata:
                # Apply metadata to table
                applied_settings = apply_table_metadata(table, metadata)

                if applied_settings:
                    print_debug_success(f"  Applied: {', '.join(applied_settings)}")
                    total_settings_applied += len(applied_settings)

                # Remove metadata from caption
                print_debug(f"  Original caption: {caption_text}")
                clean_caption = remove_metadata_from_caption(caption_text)
                caption_para.text = clean_caption
                print_debug_success("  Metadata removed from caption")

                processed_count += 1

    return processed_count, total_settings_applied


def process_file(docx_path: str, save: bool = True) -> Optional[DocumentObject]:
    """
    Process a DOCX file to apply table metadata.

    Args:
        docx_path: Path to the DOCX file
        save: Whether to save the document (default: True)

    Returns:
        Document object if successful, None otherwise
    """
    try:
        doc, docx_path_abs = open_docx(docx_path)
        if doc is None or docx_path_abs is None:
            return None

        # Process table metadata
        print_debug("Processing table metadata from captions...")
        processed, settings_applied = process_table_metadata(doc)

        print_debug_success(f"\nProcessed {processed} of {len(doc.tables)} table(s)")
        print_debug_success(f"Total settings applied: {settings_applied}")

        # Save the document
        if save:
            save_docx(doc, docx_path_abs)

        print_debug_success("\nTable metadata processing completed successfully!")
        return doc

    except Exception as e:
        print_error(f"\nTable metadata processing failed: {e}")
        import traceback
        traceback.print_exc()
        return None


def main():
    """Main entry point for command-line usage"""
    parser = argparse.ArgumentParser(
        description="Process table metadata from captions in DOCX files",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  uv run process_table_metadata.py manuscript.docx
  uv run process_table_metadata.py output/docx/manuscript.docx

Supported metadata keys:
  cell_margin=0.1cm        - Set all cell margins
  cell_margin_top=0.1cm    - Set top cell margin
  cell_margin_bottom=0.1cm - Set bottom cell margin
  cell_margin_left=0.1cm   - Set left cell margin
  cell_margin_right=0.1cm  - Set right cell margin
  row_height=1cm           - Set row height
  alignment=center         - Set table alignment (left, center, right)
  autofit=window           - Set autofit behavior (fixed, content, window)

Caption format:
  Table 1: Description |cell_margin=0.1cm alignment=center|
        """
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("--no-save", action="store_true",
                       help="Don't save the document (for testing)")

    args = parser.parse_args()

    result = process_file(args.docx_path, save=not args.no_save)
    sys.exit(0 if result else 1)


if __name__ == "__main__":
    main()
