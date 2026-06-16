#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
# ]
# ///
"""
Apply reviewer-reply visual styling to captions and table text.

This reply-only post-processor makes caption text and regular table text blue
and italic so quoted manuscript additions stand out in the response document.
Equation layout tables are skipped because they are an internal layout device,
not user-facing data tables.
"""

import argparse
import sys
from pathlib import Path
from typing import Iterable, Optional

if __package__ in (None, ""):
    # Allow direct execution via `uv run scripts/postprocess/<script>.py`.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from postprocess.common import open_docx, print_error, print_debug_success, save_docx
from postprocess.autofit_tables import is_equation_layout_table

try:
    from docx.document import Document as DocumentObject
    from docx.shared import RGBColor
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    print("Error: python-docx is not installed. Install it with: pip install python-docx")
    sys.exit(1)


REPLY_BLUE = RGBColor(0x00, 0x00, 0xFF)
CAPTION_STYLE_MARKERS = ("Caption", "题注")


def format_run_blue_italic(run) -> None:
    """Apply the reply blue italic direct formatting to a Word run."""
    run.font.color.rgb = REPLY_BLUE
    run.font.italic = True


def is_caption_paragraph(paragraph: Paragraph) -> bool:
    """Return whether a paragraph uses a caption-like style."""
    style = paragraph.style
    if style is None:
        return False
    style_name = style.name or ""
    return any(marker in style_name for marker in CAPTION_STYLE_MARKERS)


def format_caption_styles(doc: DocumentObject) -> int:
    """Make caption-like paragraph styles blue and italic."""
    updated = 0
    for style in doc.styles:
        if getattr(style, "type", None) is None:
            continue
        style_name = getattr(style, "name", "") or ""
        if any(marker in style_name for marker in CAPTION_STYLE_MARKERS):
            style.font.color.rgb = REPLY_BLUE
            style.font.italic = True
            updated += 1
    return updated


def format_caption_paragraphs(doc: DocumentObject) -> int:
    """Apply blue italic direct formatting to existing caption paragraph runs."""
    updated = 0
    for paragraph in doc.paragraphs:
        if not is_caption_paragraph(paragraph):
            continue
        for run in paragraph.runs:
            format_run_blue_italic(run)
        updated += 1
    return updated


def iter_tables_recursive(tables: Iterable[Table]) -> Iterable[Table]:
    """Yield tables and any nested tables inside their cells."""
    for table in tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from iter_tables_recursive(cell.tables)


def format_table_text(doc: DocumentObject) -> int:
    """Apply blue italic formatting to runs in regular, non-equation tables."""
    updated_runs = 0
    for table in iter_tables_recursive(doc.tables):
        if is_equation_layout_table(table):
            continue
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        format_run_blue_italic(run)
                        updated_runs += 1
    return updated_runs


def apply_reply_blue_italic_style(doc: DocumentObject) -> dict[str, int]:
    """Apply all reply-only blue italic caption and table formatting."""
    return {
        "caption_styles": format_caption_styles(doc),
        "caption_paragraphs": format_caption_paragraphs(doc),
        "table_runs": format_table_text(doc),
    }


def process_file(docx_path: str, save: bool = True) -> Optional[DocumentObject]:
    """Process a DOCX file and optionally save reply blue italic formatting."""
    try:
        doc, docx_path_abs = open_docx(docx_path)
        if doc is None or docx_path_abs is None:
            return None

        stats = apply_reply_blue_italic_style(doc)
        if save:
            save_docx(doc, docx_path_abs)

        print_debug_success(
            "Reply blue italic formatting applied: "
            f"{stats['caption_styles']} caption style(s), "
            f"{stats['caption_paragraphs']} caption paragraph(s), "
            f"{stats['table_runs']} table run(s)"
        )
        return doc
    except Exception as e:
        print_error(f"\nReply blue italic style processing failed: {e}")
        import traceback

        traceback.print_exc()
        return None


def main() -> None:
    """Main entry point for command-line usage."""
    parser = argparse.ArgumentParser(
        description="Apply blue italic formatting to reply DOCX captions and table text"
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("--no-save", action="store_true", help="Do not save changes")
    args = parser.parse_args()

    result = process_file(args.docx_path, save=not args.no_save)
    sys.exit(0 if result else 1)


if __name__ == "__main__":
    main()
