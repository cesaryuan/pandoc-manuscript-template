#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1.0",
# ]
# ///
"""
Apply the 'Where Paragraph' style to equation explanation paragraphs.

The post-processor targets paragraphs that immediately follow a display-style
math paragraph or equation layout table and begin with the word "where". This
keeps formula definitions separate from regular body paragraphs in exported
DOCX files.
"""

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import cast

if __package__ in (None, ""):
    # Allow direct execution via `uv run scripts/postprocess/<script>.py`.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from postprocess.autofit_tables import is_equation_layout_table

try:
    from docx.document import Document as DocumentObject
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.styles.style import _ParagraphStyle
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    print("Error: python-docx is not installed. Install it with: pip install python-docx")
    sys.exit(1)

from postprocess.common import (
    BODY_TEXT_STYLE_NAMES as BODY_TEXT_STYLE_CANDIDATES,
    get_first_existing_paragraph_style,
    iter_body_blocks,
    open_docx,
    print_error,
    print_debug,
    print_debug_success,
    print_warning,
    save_docx,
    set_style_first_line_indent_chars as set_common_style_first_line_indent_chars,
)


WHERE_STYLE_NAME = "Where Paragraph"
WHERE_FIRST_LINE_INDENT_CHARS = 0.0
WHERE_TABLE_LAYOUT_SPACE_BEFORE_PT = 6.0
BASE_STYLE_CANDIDATES = (*BODY_TEXT_STYLE_CANDIDATES, "First Paragraph", "Normal", "正文")
WHERE_START_PATTERN = re.compile(r"^\s*where\b", re.IGNORECASE)
EQUATION_NUMBER_PATTERN = re.compile(r"^[\s\t\r\n()（）\[\]【】0-9ivxlcdmIVXLCDM.\-–—]*$")


@dataclass
class WhereParagraphAnalysis:
    """Single-pass analysis result for where paragraphs and their equation layout mode."""

    paragraphs: list[Paragraph]
    table_layout_count: int = 0
    tab_layout_count: int = 0

    @property
    def use_table_layout_spacing(self) -> bool:
        """Return whether the Where Paragraph style should use table-layout spacing."""
        return self.table_layout_count > 0


def set_where_next_paragraph_style(doc: DocumentObject, style: _ParagraphStyle) -> bool:
    """Set the paragraph style Word uses after pressing Enter in a where paragraph."""
    body_text_style, body_text_style_name = get_first_existing_paragraph_style(
        doc, BODY_TEXT_STYLE_CANDIDATES
    )
    if body_text_style is None:
        print_error("Neither 'Body Text' nor '正文文本' style was found, cannot set next paragraph style")
        return False

    style.next_paragraph_style = body_text_style
    print_debug(f"'Where Paragraph' next paragraph style set to '{body_text_style_name}'")
    return True


def set_style_first_line_indent_chars(style: _ParagraphStyle, chars: float) -> None:
    """Set the Where Paragraph first-line indent with a contextual error label."""
    set_common_style_first_line_indent_chars(style, chars, "Where Paragraph first-line indent")


def set_where_paragraph_format(style: _ParagraphStyle, use_table_layout_spacing: bool) -> None:
    """Apply fixed paragraph formatting for the Where Paragraph style."""
    set_style_first_line_indent_chars(style, WHERE_FIRST_LINE_INDENT_CHARS)
    if use_table_layout_spacing:
        # Equation tables sit visually closer to the following where clause than tab-layout equations.
        style.paragraph_format.space_before = Pt(WHERE_TABLE_LAYOUT_SPACE_BEFORE_PT)
        print_debug(
            f"'Where Paragraph' spacing before set to {WHERE_TABLE_LAYOUT_SPACE_BEFORE_PT:g} pt"
        )
    else:
        style.paragraph_format.space_before = None
        print_debug("'Where Paragraph' spacing before inherits from base style")
    print_debug("'Where Paragraph' first-line indent set to 0 chars")


def ensure_where_paragraph_style_exists(
    doc: DocumentObject, analysis: WhereParagraphAnalysis | None = None
) -> bool:
    """Ensure the DOCX contains a paragraph style named 'Where Paragraph'."""
    analysis = analysis or analyze_where_paragraphs(doc)
    log_where_layout_detection(analysis)
    try:
        style = cast(_ParagraphStyle, doc.styles[WHERE_STYLE_NAME])
        if not set_where_next_paragraph_style(doc, style):
            return False
        set_where_paragraph_format(style, analysis.use_table_layout_spacing)
        print_debug("'Where Paragraph' style already exists")
        return True
    except KeyError:
        print_debug("'Where Paragraph' style not found, creating it...")

    try:
        style = cast(_ParagraphStyle, doc.styles.add_style(WHERE_STYLE_NAME, WD_STYLE_TYPE.PARAGRAPH))
        base_style, base_style_name = get_first_existing_paragraph_style(doc, BASE_STYLE_CANDIDATES)
        if base_style is not None:
            style.base_style = base_style
            print_debug(f"'Where Paragraph' style based on '{base_style_name}' style")
        if not set_where_next_paragraph_style(doc, style):
            return False
        set_where_paragraph_format(style, analysis.use_table_layout_spacing)

        # Where clauses are equation definitions, so keep them available as a named style.
        style.hidden = False
        style.quick_style = True
        style.priority = 1
        print_debug_success("'Where Paragraph' style created successfully")
        return True
    except Exception as e:
        print_error(f"Failed to create 'Where Paragraph' style: {e}")
        return False


def paragraph_contains_math(paragraph: Paragraph) -> bool:
    """Return whether a paragraph contains Word math elements."""
    element = paragraph._p
    return bool(element.findall(f".//{qn('m:oMath')}") or element.findall(f".//{qn('m:oMathPara')}"))


def paragraph_has_equation_layout(paragraph: Paragraph) -> bool:
    """Return whether a math paragraph looks like a standalone/display equation."""
    element = paragraph._p
    return bool(element.findall(f".//{qn('w:tab')}") or element.findall(f".//{qn('w:tabs')}"))


def paragraph_visible_text(paragraph: Paragraph) -> str:
    """Return visible non-math text from a paragraph."""
    return paragraph.text or ""


def is_equation_paragraph(paragraph: Paragraph) -> bool:
    """Return whether a paragraph should be treated as the formula before a where clause."""
    if not paragraph_contains_math(paragraph):
        return False
    if paragraph_has_equation_layout(paragraph):
        return True

    # Unnumbered display equations may contain only math and punctuation visible to python-docx.
    return bool(EQUATION_NUMBER_PATTERN.match(paragraph_visible_text(paragraph)))


def starts_with_where(paragraph: Paragraph) -> bool:
    """Return whether a paragraph begins with the word 'where'."""
    return bool(WHERE_START_PATTERN.match(paragraph_visible_text(paragraph)))


def is_equation_block(block: Paragraph | Table) -> bool:
    """Return whether a body block is an equation paragraph or equation layout table."""
    if isinstance(block, Paragraph):
        return is_equation_paragraph(block)
    return is_equation_layout_table(block)


def analyze_where_paragraphs(doc: DocumentObject) -> WhereParagraphAnalysis:
    """Find where paragraphs and detect whether their preceding equations use tables or tabs."""
    analysis = WhereParagraphAnalysis(paragraphs=[])
    previous: Paragraph | Table | None = None

    for current in iter_body_blocks(doc):
        if not isinstance(current, Paragraph):
            previous = current
            continue

        if previous is None or not starts_with_where(current):
            previous = current
            continue

        if isinstance(previous, Table) and is_equation_layout_table(previous):
            analysis.paragraphs.append(current)
            analysis.table_layout_count += 1
        elif isinstance(previous, Paragraph) and is_equation_paragraph(previous):
            analysis.paragraphs.append(current)
            if paragraph_has_equation_layout(previous):
                analysis.tab_layout_count += 1

        previous = current

    return analysis


def log_where_layout_detection(analysis: WhereParagraphAnalysis) -> None:
    """Log the detected equation layout mode used before where paragraphs."""
    if analysis.table_layout_count:
        print_debug(
            f"Detected {analysis.table_layout_count} where paragraph(s) after equation layout table(s); "
            f"using {WHERE_TABLE_LAYOUT_SPACE_BEFORE_PT:g} pt style spacing before"
        )
        return

    if analysis.tab_layout_count:
        print_debug(
            f"Detected {analysis.tab_layout_count} where paragraph(s) after tab-layout equation paragraph(s); "
            "leaving style spacing before inherited"
        )
    else:
        print_debug("No table-layout where paragraphs detected; leaving style spacing before inherited")


def apply_where_paragraph_style(
    doc: DocumentObject, analysis: WhereParagraphAnalysis | None = None
) -> int:
    """Style where paragraphs that immediately follow equation paragraphs or tables."""
    analysis = analysis or analyze_where_paragraphs(doc)
    for paragraph in analysis.paragraphs:
        paragraph.style = WHERE_STYLE_NAME
    return len(analysis.paragraphs)


def process_where_paragraph_styles(doc: DocumentObject) -> int | None:
    """Analyze once, configure the Where Paragraph style, and apply it to matching paragraphs."""
    analysis = analyze_where_paragraphs(doc)
    if not ensure_where_paragraph_style_exists(doc, analysis):
        return None
    return apply_where_paragraph_style(doc, analysis)


def process_file(docx_path: str, save: bool = True) -> int | None:
    """Process a DOCX file and optionally save where paragraph style changes."""
    doc, docx_path_abs = open_docx(docx_path)
    if doc is None or docx_path_abs is None:
        return None
    updated = process_where_paragraph_styles(doc)
    if updated is None:
        return None
    if save:
        save_docx(doc, docx_path_abs)

    print_debug_success(f"Applied 'Where Paragraph' style to {updated} paragraph(s)")
    return updated


def main() -> None:
    """Main entry point for command-line usage."""
    parser = argparse.ArgumentParser(
        description="Apply 'Where Paragraph' style after DOCX equations"
    )
    parser.add_argument("docx_path", help="Path to the DOCX file to process")
    parser.add_argument("--no-save", action="store_true", help="Do not save changes")
    args = parser.parse_args()

    result = process_file(args.docx_path, save=not args.no_save)
    sys.exit(0 if result is not None else 1)


if __name__ == "__main__":
    main()
