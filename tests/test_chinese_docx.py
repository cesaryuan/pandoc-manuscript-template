"""Verify user-visible Chinese DOCX builds and style customization."""

from pathlib import Path
import os
import shutil
import subprocess
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn


PNG_PIXEL = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000b49444154789c636000020000050001a5f645400000000049454e44ae426082"
)


def run_docx_build(project: Path, cli_args: list[str]) -> tuple[str, Path]:
    """Run the public build command in an isolated manuscript project."""
    output = project / "result.docx"
    environment = os.environ.copy()
    environment["LANG"] = "en-US"
    result = subprocess.run(
        [sys.executable, "-m", "pandoc_manuscript.cli", "build", "docx", "paper.md", "--no-mathtype", "-o", str(output), *cli_args],
        cwd=project,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        env=environment,
        timeout=120,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return result.stdout + result.stderr, output


@pytest.mark.parametrize(
    ("metadata_lang", "style_lang", "cli_args"),
    [
        ("zh-CN", None, []),
        (None, "zh-CN", []),
        ("en-US", None, ["--lang", "zh-cn"]),
    ],
)
def test_chinese_docx_build_numbers_figures_tables_and_formats_headings(
    tmp_path: Path,
    metadata_lang: str | None,
    style_lang: str | None,
    cli_args: list[str],
) -> None:
    """Use manuscript/style metadata or CLI override for the same Chinese DOCX behavior."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "figure.png").write_bytes(PNG_PIXEL)
    manuscript = tmp_path / "paper.md"
    original = (
        "---\n"
        + (f"lang: {metadata_lang}\n" if metadata_lang is not None else "")
        + "title: 中文标题\n"
        "figureTitle: 图\ntableTitle: 表\ntitleDelim: ' '\n---\n\n"
        "# 第一章\n\n# 第二章\n\n# 第三章\n\n"
        "![测试图](figure.png){#fig:one}\n\n"
        "| 列一 | 列二 |\n| --- | --- |\n| 值一 | 值二 |\n\n: 测试表 {#tbl:one}\n"
    )
    manuscript.write_text(original, encoding="utf-8")
    if style_lang is not None:
        (tmp_path / "style.yml").write_text(
            f"pandocMetadata:\n  lang: {style_lang}\n",
            encoding="utf-8",
        )

    output_log, output = run_docx_build(tmp_path, cli_args)

    assert "Could not load translations" not in output_log
    assert "has no translation defined" not in output_log
    assert manuscript.read_text(encoding="utf-8") == original
    assert not list(tmp_path.glob(".paper.pmt-no-lang-*.md"))
    doc = Document(output)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "图 3-1" in text
    assert "表 3-1" in text
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        assert style.font.bold is False, name
        assert style.element.rPr.bCs.val is False, name
        fonts = style.element.rPr.rFonts
        assert fonts.get(qn("w:eastAsia")) == "黑体", name
        assert fonts.get(qn("w:ascii")) == "黑体", name
        assert fonts.get(qn("w:eastAsiaTheme")) is None, name


def test_docx_style_font_and_bold_work_without_language_mode(tmp_path: Path) -> None:
    """Apply explicitly configured font and weight even for a non-Chinese build."""
    if not shutil.which("pandoc") or not shutil.which("pandoc-crossref"):
        pytest.skip("Pandoc and pandoc-crossref are required for the DOCX contract")
    (tmp_path / "paper.md").write_text("# Heading\n\nBody.\n", encoding="utf-8")
    (tmp_path / "style.yml").write_text(
        "docxStyle:\n  标题 1: {fontFamily: 黑体, bold: false}\n",
        encoding="utf-8",
    )

    _, output = run_docx_build(tmp_path, [])

    style = Document(output).styles["Heading 1"]
    assert style.font.bold is False
    assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"
