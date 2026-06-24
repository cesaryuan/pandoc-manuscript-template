from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from pandoc_manuscript.docx.postprocess.process_equation_metadata import process_equation_metadata


def test_equation_revision_attr_filter_wraps_display_equation_and_keeps_label() -> None:
    """Preserve the equation label while extracting revision=true before crossref."""
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        pytest.skip("pandoc is not installed")

    filter_path = Path(__file__).resolve().parents[1] / "pandoc" / "filters" / "equation_revision_attr.lua"
    markdown = """\
$$
a+b
$$ {#eq:sum revision=true}
"""

    result = subprocess.run(
        [pandoc, "--lua-filter", str(filter_path), "-f", "markdown", "-t", "native"],
        input=markdown,
        text=True,
        capture_output=True,
        check=True,
    )

    assert '( "revision" , "true" )' in result.stdout
    assert "{#eq:sum}" in result.stdout


def test_docx_metadata_filter_emits_equation_revision_marker(tmp_path) -> None:
    """Export revised display equations through the shared DOCX metadata filter."""
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        pytest.skip("pandoc is not installed")

    filter_path = Path(__file__).resolve().parents[1] / "pandoc" / "filters" / "docx_metadata.lua"
    output_path = tmp_path / "equation-metadata.docx"
    markdown = """\
::: {revision=true}
$$
a+b
$$
:::
"""

    subprocess.run(
        [pandoc, "--lua-filter", str(filter_path), "-f", "markdown", "-o", str(output_path)],
        input=markdown,
        text=True,
        capture_output=True,
        check=True,
    )

    with zipfile.ZipFile(output_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "PMT_EQUATION_METADATA:" in document_xml


def test_process_equation_metadata_colors_native_word_display_equations(tmp_path) -> None:
    """Color native Word display-equation runs red when the hidden marker is present."""
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        pytest.skip("pandoc is not installed")

    markdown_path = tmp_path / "equation.md"
    markdown_path.write_text("$$\na+b\n$$\n", encoding="utf-8")
    docx_path = tmp_path / "equation.docx"

    subprocess.run(
        [pandoc, str(markdown_path), "-o", str(docx_path)],
        text=True,
        capture_output=True,
        check=True,
    )

    doc = Document(str(docx_path))
    equation_paragraph = next(
        paragraph for paragraph in doc.paragraphs if paragraph._p.findall(f".//{qn('m:oMathPara')}")
    )
    equation_paragraph.insert_paragraph_before('PMT_EQUATION_METADATA:{"revision":"true"}')

    processed, updated_runs = process_equation_metadata(doc)

    assert processed == 1
    assert updated_runs > 0
    assert all("PMT_EQUATION_METADATA:" not in paragraph.text for paragraph in doc.paragraphs)

    math_runs = equation_paragraph._p.findall(f".//{qn('m:r')}")
    assert math_runs
    for math_run in math_runs:
        math_rpr = math_run.find(qn("m:rPr"))
        assert math_rpr is not None
        word_rpr = math_rpr.find(qn("w:rPr"))
        assert word_rpr is not None
        color = word_rpr.find(qn("w:color"))
        assert color is not None
        assert color.get(qn("w:val")) == "FF0000"
