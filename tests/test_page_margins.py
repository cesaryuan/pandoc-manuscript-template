from pathlib import Path
import sys

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from pandoc_manuscript.commands import build
from pandoc_manuscript.commands.build_reply import output as reply_output
from pandoc_manuscript.docx.page_margins import (
    apply_page_margin_metadata,
    normalize_page_margins,
    write_reference_doc_with_page_margins,
)


def test_apply_page_margin_metadata_updates_all_docx_sections() -> None:
    """Apply docxPageMargins values to every generated DOCX section."""
    doc = Document()
    doc.add_section()

    result = apply_page_margin_metadata(
        doc,
        {
            "docxPageMargins": {
                "top": "2.54cm",
                "bottom": "2.54cm",
                "left": "3.17cm",
                "right": "3.17cm",
            }
        },
    )

    assert result == {
        "margins": {
            "top": "2.54cm",
            "bottom": "2.54cm",
            "left": "3.17cm",
            "right": "3.17cm",
        },
        "sections": 2,
    }
    for section in doc.sections:
        assert section.top_margin.cm == pytest.approx(2.54, abs=0.001)
        assert section.bottom_margin.cm == pytest.approx(2.54, abs=0.001)
        assert section.left_margin.cm == pytest.approx(3.17, abs=0.001)
        assert section.right_margin.cm == pytest.approx(3.17, abs=0.001)


def test_docx_page_margins_partial_config_leaves_other_sides_unset() -> None:
    """Do not invent margin defaults when only some sides are configured."""
    normalized = normalize_page_margins({"docxPageMargins": {"top": "72pt"}})

    assert normalized is not None
    margins, display_values = normalized
    assert set(margins) == {"top"}
    assert display_values == {"top": "72pt"}


def test_write_reference_doc_with_page_margins_updates_copy(tmp_path: Path) -> None:
    """Apply docxPageMargins to the reference DOCX copy Pandoc will read."""
    source = tmp_path / "source.docx"
    target = tmp_path / "target.docx"
    Document().save(source)

    result = write_reference_doc_with_page_margins(
        source,
        target,
        {
            "docxPageMargins": {
                "top": "2.54cm",
                "bottom": "2.54cm",
                "left": "3.17cm",
                "right": "3.17cm",
            }
        },
    )

    assert result is not None
    assert result["target"] == target
    section = Document(str(target)).sections[0]
    assert section.top_margin.cm == pytest.approx(2.54, abs=0.001)
    assert section.bottom_margin.cm == pytest.approx(2.54, abs=0.001)
    assert section.left_margin.cm == pytest.approx(3.17, abs=0.001)
    assert section.right_margin.cm == pytest.approx(3.17, abs=0.001)


def test_build_reference_doc_args_uses_margin_adjusted_reference(tmp_path: Path, monkeypatch) -> None:
    """Pass Pandoc a generated reference DOCX when docxPageMargins is configured."""
    source = tmp_path / "reference.docx"
    Document().save(source)
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(build.SETTINGS, "reference_doc", str(source))
    monkeypatch.setattr(build.SETTINGS, "project_name", "paper")

    args = build.docx_reference_doc_args({"docxPageMargins": {"left": "3.17cm", "right": "3.17cm"}})

    assert args[0] == "--reference-doc"
    assert Path(args[1]).parts[:3] == (".pmt", "work", "reference-doc")
    section = Document(args[1]).sections[0]
    assert section.left_margin.cm == pytest.approx(3.17, abs=0.001)
    assert section.right_margin.cm == pytest.approx(3.17, abs=0.001)


def test_reply_reference_doc_for_pandoc_uses_margin_adjusted_reference(tmp_path: Path, monkeypatch) -> None:
    """Prepare a generated reply reference DOCX when docxPageMargins is configured."""
    source = tmp_path / "reference.docx"
    output = tmp_path / "reply.docx"
    Document().save(source)
    monkeypatch.chdir(tmp_path)

    reference = reply_output.reply_reference_doc_for_pandoc(
        source,
        output,
        {"docxPageMargins": {"left": "3.17cm", "right": "3.17cm"}},
    )

    assert reference.parts[:4] == (".pmt", "work", "reply", "reference-doc")
    section = Document(str(reference)).sections[0]
    assert section.left_margin.cm == pytest.approx(3.17, abs=0.001)
    assert section.right_margin.cm == pytest.approx(3.17, abs=0.001)
