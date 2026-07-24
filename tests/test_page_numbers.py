from pathlib import Path
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from pandoc_manuscript.docx.postprocess.page_numbers import (
    PAGE_NUMBER_STYLE_NAME,
    apply_page_number_settings,
)
from pandoc_manuscript.runtime.metadata import PmtSettings


def footer_page_instructions(doc: Document) -> list[str]:
    """Return Word field instructions from the primary footer of each section."""
    return [
        instruction.text or ""
        for section in doc.sections
        for paragraph in section.footer.paragraphs
        for instruction in paragraph._p.iter(qn("w:instrText"))
    ]


def add_page_number_style(doc: Document) -> None:
    """Add the required reference-DOCX page number style to a test document."""
    doc.styles.add_style(PAGE_NUMBER_STYLE_NAME, WD_STYLE_TYPE.CHARACTER)


def test_docx_page_numbers_are_unchanged_when_not_configured() -> None:
    """Preserve custom reference-DOCX footer behavior until the setting is explicit."""
    doc = Document()

    assert apply_page_number_settings(doc, PmtSettings.model_validate({})) is None
    assert footer_page_instructions(doc) == []


def test_docx_page_numbers_can_be_added_and_removed(tmp_path: Path) -> None:
    """Insert a PAGE field for true and remove it again for false."""
    doc = Document()
    add_page_number_style(doc)
    doc.add_section()

    result = apply_page_number_settings(
        doc,
        PmtSettings.model_validate({"docxShowPageNumbers": True}),
    )

    assert result == {"visible": True, "added": 1, "existing": 0}
    assert footer_page_instructions(doc) == [" PAGE ", " PAGE "]
    assert all(
        run.style.name == PAGE_NUMBER_STYLE_NAME
        for run in doc.sections[0].footer.paragraphs[0].runs
    )
    assert doc.settings.element.find(qn("w:updateFields")) is not None

    output = tmp_path / "page-numbers.docx"
    doc.save(output)
    doc = Document(output)
    assert footer_page_instructions(doc) == [" PAGE ", " PAGE "]
    assert all(
        run.style.name == PAGE_NUMBER_STYLE_NAME
        for run in doc.sections[0].footer.paragraphs[0].runs
    )

    result = apply_page_number_settings(
        doc,
        PmtSettings.model_validate({"docxShowPageNumbers": False}),
    )

    assert result == {"visible": False, "removed": 1}
    assert footer_page_instructions(doc) == []
