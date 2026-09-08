"""Check saved DOCX equation styling without changing unrelated paragraph layouts."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from pandoc_manuscript.docx.postprocess.para_equation_style import process_para_equation_style


def add_math(paragraph, ole: bool = False) -> None:
    """Insert an OMML formula or a MathType object marker for layout recognition."""
    if ole:
        obj = OxmlElement("w:object")
        etree.SubElement(obj, "{urn:schemas-microsoft-com:office:office}OLEObject", ProgID="Equation.DSMT4")
        paragraph.add_run()._r.append(obj)
    else:
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = "x=1"
        run.append(text)
        math.append(run)
        paragraph._p.append(math)


@pytest.mark.parametrize("ole", [False, True])
def test_saved_tab_equation_inherits_body_style_without_direct_spacing_overrides(tmp_path, ole) -> None:
    """Preserve content while migrating direct tabs and spacing into the equation style."""
    doc = Document()
    doc.sections[0].page_width = Inches(8.5)
    doc.sections[0].left_margin = Inches(1)
    doc.sections[0].right_margin = Inches(1)
    doc.styles["Body Text"].font.size = Pt(13)
    old_style = doc.styles.add_style("Para Equation", WD_STYLE_TYPE.PARAGRAPH)
    old_style.paragraph_format.space_after = Pt(24)
    equation = doc.add_paragraph("\t")
    equation.alignment = WD_ALIGN_PARAGRAPH.LEFT
    text_alignment = OxmlElement("w:textAlignment")
    text_alignment.set(qn("w:val"), "top")
    equation._p.get_or_add_pPr().append(text_alignment)
    add_math(equation, ole=ole)
    equation.add_run("\t(1)")
    equation.paragraph_format.tab_stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.CENTER)
    equation.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
    equation.paragraph_format.space_before = Pt(3)
    equation.paragraph_format.space_after = Pt(18)
    equation.paragraph_format.line_spacing = 2.0
    spacing = equation._p.pPr.find(qn("w:spacing"))
    spacing.set(qn("w:afterLines"), "200")
    spacing.set(qn("w:afterAutospacing"), "1")
    original_content = [etree.tostring(child) for child in equation._p if child.tag != qn("w:pPr")]

    assert process_para_equation_style(doc) == 1
    path = tmp_path / "equation.docx"
    doc.save(path)
    reopened = Document(path)
    paragraph = reopened.paragraphs[0]
    assert paragraph.style.name == "Para Equation"
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraph.style.paragraph_format.alignment is None
    assert paragraph._p.pPr.find(qn("w:textAlignment")) is None
    assert paragraph.style.element.pPr.find(qn("w:textAlignment")).get(qn("w:val")) == "center"
    assert paragraph.style.base_style.name == "Body Text"
    assert paragraph.style.base_style.font.size.pt == 13
    assert paragraph.style.paragraph_format.line_spacing == 1.0
    style_spacing = paragraph.style.element.pPr.find(qn("w:spacing"))
    assert style_spacing.get(qn("w:afterLines")) == "50"
    assert style_spacing.get(qn("w:afterAutospacing")) == "0"
    assert style_spacing.get(qn("w:after")) is None
    assert paragraph.paragraph_format.space_before.pt == 3
    assert dict(paragraph._p.pPr.find(qn("w:spacing")).attrib) == {qn("w:before"): "60"}
    assert paragraph._p.pPr.find(qn("w:tabs")) is None
    assert [(tab.position.inches, tab.alignment) for tab in paragraph.style.paragraph_format.tab_stops] == [
        (3.25, WD_TAB_ALIGNMENT.CENTER), (6.5, WD_TAB_ALIGNMENT.RIGHT),
    ]
    assert [etree.tostring(child) for child in paragraph._p if child.tag != qn("w:pPr")] == original_content
    first_pass = etree.tostring(reopened.element)
    assert process_para_equation_style(reopened) == 1
    assert etree.tostring(reopened.element) == first_pass


def test_equation_style_tabs_follow_page_changes_and_cancel_inherited_stops(tmp_path) -> None:
    """Recompute saved style tabs after a margin change without inherited tab interference."""
    doc = Document()
    doc.sections[0].page_width = Inches(8.5)
    doc.sections[0].left_margin = Inches(1)
    doc.sections[0].right_margin = Inches(1)
    doc.styles["Body Text"].paragraph_format.tab_stops.add_tab_stop(Inches(1))
    equation = doc.add_paragraph()
    add_math(equation)
    # Even a template with only tab definitions must remain recognizable after migration.
    equation.paragraph_format.tab_stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.CENTER)
    assert process_para_equation_style(doc) == 1
    doc.sections[0].right_margin = Inches(1.5)
    assert process_para_equation_style(doc) == 1
    path = tmp_path / "resized.docx"
    doc.save(path)
    reopened = Document(path)
    tabs = reopened.styles["Para Equation"].paragraph_format.tab_stops
    assert [(tab.position.inches, tab.alignment) for tab in tabs] == [
        (1, WD_TAB_ALIGNMENT.CLEAR),
        (3, WD_TAB_ALIGNMENT.CENTER),
        (6, WD_TAB_ALIGNMENT.RIGHT),
    ]
    assert reopened.paragraphs[0]._p.pPr.find(qn("w:tabs")) is None
    assert len(reopened.styles["Body Text"].paragraph_format.tab_stops) == 1


def test_non_tab_math_and_table_paragraphs_are_untouched() -> None:
    """Exclude ordinary tab text, inline/display math without tabs, and table equations."""
    doc = Document()
    doc.add_paragraph("Label\tValue")
    add_math(doc.add_paragraph("Inline formula: "))
    add_math(doc.add_paragraph())
    cell_paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    cell_paragraph.add_run("\t")
    add_math(cell_paragraph)
    before = etree.tostring(doc.element)
    assert process_para_equation_style(doc) == 0
    assert etree.tostring(doc.element) == before
    assert "Para Equation" not in doc.styles
